from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def set_run_font(run, size: int, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)


def set_paragraph_font(paragraph, size: int) -> None:
    for run in paragraph.runs:
      run.font.name = "Times New Roman"
      run.font.size = Pt(size)


def add_centered_line(document: Document, text: str, size: int, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size, bold=bold)


def add_left_line(document: Document, text: str, size: int, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_run_font(run, size, bold=bold)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_run_font(run, 11)


def add_single_image(document: Document, image_path: Path, max_width_inches: float = 6.5) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(max_width_inches))


def clear_table_borders(table) -> None:
    table_pr = table._tbl.tblPr
    for child in list(table_pr):
        if child.tag.endswith("tblBorders"):
            table_pr.remove(child)


def add_image_grid(document: Document, image_paths: list[Path], columns: int = 2, width_inches: float = 3.15) -> None:
    rows = (len(image_paths) + columns - 1) // columns
    table = document.add_table(rows=rows, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    clear_table_borders(table)

    for index, image_path in enumerate(image_paths):
        cell = table.cell(index // columns, index % columns)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Inches(width_inches))


def load_manifest(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def build_document(manifest: dict, group_members_text: str) -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    header_paragraph = section.header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_paragraph.add_run(manifest["header"])
    set_run_font(header_run, 10)

    add_centered_line(document, manifest["university"], 12)
    add_centered_line(document, manifest["course"], 12)
    document.add_paragraph()
    add_centered_line(document, manifest["titleLabel"], 12, bold=True)
    add_centered_line(document, manifest["projectTitle"], 12, bold=True)
    document.add_paragraph()
    add_centered_line(document, group_members_text or manifest["groupMembersLabel"], 12)
    document.add_paragraph()

    add_left_line(document, manifest["transactionModuleTitle"], 12, bold=True)

    for section_data in manifest["sections"]:
        add_left_line(document, section_data["title"], 12, bold=True)
        add_bullet(document, section_data["description"])

        image_paths = [Path(item["path"]) for item in section_data["images"]]
        if not image_paths:
            continue

        if section_data.get("layout") == "single" and len(image_paths) == 1:
            add_single_image(document, image_paths[0])
        else:
            add_image_grid(document, image_paths)

        document.add_paragraph()

    return document


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--group-members", default="")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    manifest = load_manifest(Path(args.manifest))
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    document = build_document(manifest, args.group_members.strip())
    document.save(output_path)


if __name__ == "__main__":
    main()
